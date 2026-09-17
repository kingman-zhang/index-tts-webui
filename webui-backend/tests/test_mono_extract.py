"""mono 文档导入（/api/mono/extract 逻辑层）单测。

运行：webui-backend 目录下
  /Users/zhangjianwen/.workbuddy/binaries/python/envs/default/bin/python tests/test_mono_extract.py
"""
import io
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi import HTTPException

from app.routes.mono import (
    _decode_text,
    _extract_docx,
    _extract_pdf,
    IMPORT_MAX_CHARS,
)

PASS = 0
FAIL = 0


def check(cond, msg):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"PASS {msg}")
    else:
        FAIL += 1
        print(f"FAIL {msg}")


def expect_http_error(fn, status, msg):
    global PASS, FAIL
    try:
        fn()
        FAIL += 1
        print(f"FAIL {msg}（未抛出 HTTPException）")
    except HTTPException as e:
        ok = e.status_code == status
        if ok:
            PASS += 1
            print(f"PASS {msg}")
        else:
            FAIL += 1
            print(f"FAIL {msg}（status={e.status_code}）")
    except Exception as e:
        FAIL += 1
        print(f"FAIL {msg}（异常类型 {type(e).__name__}: {e}）")


# ── 1. txt/md 解码 ───────────────────────────────────────────
check(_decode_text("用来充实我们的大脑".encode("utf-8")) == "用来充实我们的大脑", "txt utf-8")
check(_decode_text("恐惧情绪测试".encode("gb18030")) == "恐惧情绪测试", "txt gb18030")
check(_decode_text(b"\xef\xbb\xbfBOM text") == "BOM text", "txt utf-8-sig BOM")

# ── 2. docx 提取 ─────────────────────────────────────────────
try:
    import docx as docx_mod

    doc = docx_mod.Document()
    doc.add_paragraph("第一段：读书。")
    doc.add_paragraph("")  # 空段应被丢弃
    doc.add_paragraph("第二段：看电影；")
    buf = io.BytesIO()
    doc.save(buf)
    text = _extract_docx(buf.getvalue())
    check(text == "第一段：读书。\n第二段：看电影；", f"docx 段落提取（空段丢弃）: {text!r}")
except ImportError:
    print("SKIP docx 测试（python-docx 未安装）")

# ── 3. pdf 页数上限 ──────────────────────────────────────────
try:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(51):
        writer.add_blank_page(width=100, height=100)
    buf = io.BytesIO()
    writer.write(buf)
    expect_http_error(lambda: _extract_pdf(buf.getvalue()), 400, "pdf 51 页拒绝")
except ImportError:
    print("SKIP pdf 测试（pypdf 未安装）")

# ── 4. 字数统计口径（空白不计） ──────────────────────────────
sample = "用 来\n充实大脑\t[pause:0.5]\n" * 3
chars = len(re.sub(r"\s", "", sample))
check(chars == len("用来充实大脑[pause:0.5]") * 3, f"字数统计不含空白: {chars}")
check(IMPORT_MAX_CHARS == 10000, "字数上限 1 万")

print(f"\n{PASS} passed, {FAIL} failed")
sys.exit(1 if FAIL else 0)
